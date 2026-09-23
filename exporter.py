import re
from pathlib import Path


# ==================== Markdown 解析 ====================
def parse_markdown(text: str) -> list[dict]:
    """
    把 Markdown 解析成块列表：
    [
        {"type": "h1", "text": "标题"},
        {"type": "p",  "text": "段落"},
        {"type": "bullet", "text": "列表项"},
        {"type": "quote", "text": "引用"},
        {"type": "code", "text": "代码"},
        {"type": "hr"},
    ]
    """
    blocks = []
    in_code = False
    code_buf = []

    for raw in text.splitlines():
        line = raw.rstrip()

        # 代码块
        if line.startswith("```"):
            if in_code:
                blocks.append({"type": "code", "text": "\n".join(code_buf)})
                code_buf = []
                in_code = False
            else:
                in_code = True
            continue
        if in_code:
            code_buf.append(line)
            continue

        if not line.strip():
            continue

        # 分割线
        if re.match(r"^(-{3,}|\*{3,}|_{3,})$", line.strip()):
            blocks.append({"type": "hr"})
            continue

        # 标题
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            level = len(m.group(1))
            blocks.append({"type": f"h{level}", "text": m.group(2).strip()})
            continue

        # 引用
        if line.startswith(">"):
            blocks.append({"type": "quote", "text": line.lstrip("> ").strip()})
            continue

        # 列表
        m = re.match(r"^\s*[-*+]\s+(.*)$", line)
        if m:
            blocks.append({"type": "bullet", "text": m.group(1).strip()})
            continue

        m = re.match(r"^\s*\d+\.\s+(.*)$", line)
        if m:
            blocks.append({"type": "number", "text": m.group(1).strip()})
            continue

        # 普通段落
        blocks.append({"type": "p", "text": line.strip()})

    if in_code and code_buf:
        blocks.append({"type": "code", "text": "\n".join(code_buf)})

    return blocks


# ==================== Markdown ====================
def export_markdown(text: str, path: Path, title: str = ""):
    content = text
    if title and not text.lstrip().startswith("# "):
        content = f"# {title}\n\n{text}"
    path.write_text(content, encoding="utf-8")


# ==================== Word ====================
def export_docx(text: str, path: Path, title: str = "笔记"):
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # 文档标题
    if title:
        h = doc.add_heading(title, level=0)
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER

    blocks = parse_markdown(text)
    for b in blocks:
        t = b["type"]

        if t == "h1":
            doc.add_heading(b["text"], level=1)
        elif t == "h2":
            doc.add_heading(b["text"], level=2)
        elif t == "h3":
            doc.add_heading(b["text"], level=3)
        elif t in ("h4", "h5", "h6"):
            doc.add_heading(b["text"], level=4)
        elif t == "bullet":
            doc.add_paragraph(b["text"], style="List Bullet")
        elif t == "number":
            doc.add_paragraph(b["text"], style="List Number")
        elif t == "quote":
            p = doc.add_paragraph()
            run = p.add_run(b["text"])
            run.italic = True
            run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
        elif t == "code":
            p = doc.add_paragraph()
            run = p.add_run(b["text"])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        elif t == "hr":
            doc.add_paragraph("─" * 30)
        else:
            doc.add_paragraph(b["text"])

    doc.save(str(path))


# ==================== PDF ====================
def export_pdf(text: str, path: Path, title: str = "笔记"):
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.lib import colors
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, HRFlowable, Preformatted,
    )
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont

    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))

    styles = getSampleStyleSheet()
    base = ParagraphStyle(
        "Base", parent=styles["Normal"],
        fontName="STSong-Light", fontSize=11, leading=18,
    )
    h1 = ParagraphStyle("H1", parent=base, fontSize=20, leading=26,
                        spaceAfter=12, textColor=colors.HexColor("#1a1a1a"))
    h2 = ParagraphStyle("H2", parent=base, fontSize=16, leading=22,
                        spaceAfter=8, textColor=colors.HexColor("#2a2a2a"))
    h3 = ParagraphStyle("H3", parent=base, fontSize=13, leading=20,
                        spaceAfter=6, textColor=colors.HexColor("#3a3a3a"))
    quote = ParagraphStyle("Quote", parent=base, fontSize=10.5, leading=18,
                           leftIndent=12, textColor=colors.HexColor("#666666"))
    code = ParagraphStyle("Code", parent=base, fontName="STSong-Light",
                          fontSize=9.5, leading=14,
                          backColor=colors.HexColor("#f4f4f4"),
                          borderPadding=6)

    doc = SimpleDocTemplate(
        str(path), pagesize=A4,
        leftMargin=2 * cm, rightMargin=2 * cm,
        topMargin=2 * cm, bottomMargin=2 * cm,
        title=title or "笔记",
    )

    story = []
    if title:
        story.append(Paragraph(title, h1))
        story.append(Spacer(1, 6))

    blocks = parse_markdown(text)
    for b in blocks:
        t = b["type"]
        if t == "h1":
            story.append(Paragraph(b["text"], h1))
        elif t == "h2":
            story.append(Paragraph(b["text"], h2))
        elif t in ("h3", "h4", "h5", "h6"):
            story.append(Paragraph(b["text"], h3))
        elif t == "bullet":
            story.append(Paragraph(f"• {b['text']}", base))
        elif t == "number":
            story.append(Paragraph(f"‣ {b['text']}", base))
        elif t == "quote":
            story.append(Paragraph(f"「{b['text']}」", quote))
        elif t == "code":
            story.append(Preformatted(b["text"], code))
            story.append(Spacer(1, 6))
        elif t == "hr":
            story.append(Spacer(1, 4))
            story.append(HRFlowable(width="100%", color=colors.HexColor("#cccccc")))
            story.append(Spacer(1, 6))
        else:
            # 粗体 **xxx**
            html = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", b["text"])
            story.append(Paragraph(html, base))

    doc.build(story)