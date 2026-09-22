from pathlib import Path


# ==================== Markdown → Word ====================
def md_to_docx(md_path: Path, out_path: Path):
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from exporter import parse_markdown

    text = md_path.read_text(encoding="utf-8")
    doc = Document()

    for b in parse_markdown(text):
        t = b["type"]
        if t == "h1":
            doc.add_heading(b["text"], level=1)
        elif t == "h2":
            doc.add_heading(b["text"], level=2)
        elif t == "h3":
            doc.add_heading(b["text"], level=3)
        elif t in ("h4", "h5", "h6"):
            doc.add_heading(b["text"], level=4)
        elif t == "bullet":
            doc.add_paragraph(b["text"], style="List Bullet")
        elif t == "number":
            doc.add_paragraph(b["text"], style="List Number")
        elif t == "quote":
            p = doc.add_paragraph()
            run = p.add_run(b["text"])
            run.italic = True
            run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)
        elif t == "code":
            p = doc.add_paragraph()
            run = p.add_run(b["text"])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        elif t == "hr":
            doc.add_paragraph("─" * 30)
        else:
            doc.add_paragraph(b["text"])

    doc.save(str(out_path))


# ==================== Word → Markdown ====================
def docx_to_md(docx_path: Path, out_path: Path):
    from docx import Document

    doc = Document(str(docx_path))
    lines = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            lines.append("")
            continue

        style = (para.style.name or "").lower()
        if style.startswith("heading 1") or style == "title":
            lines.append(f"# {text}")
        elif style.startswith("heading 2"):
            lines.append(f"## {text}")
        elif style.startswith("heading 3"):
            lines.append(f"### {text}")
        elif style.startswith("heading"):
            lines.append(f"#### {text}")
        elif "list bullet" in style:
            lines.append(f"- {text}")
        elif "list number" in style:
            lines.append(f"1. {text}")
        else:
            lines.append(text)

    out_path.write_text("\n".join(lines), encoding="utf-8")


# ==================== 图片 → PDF ====================
def images_to_pdf(image_paths: list[Path], out_path: Path):
    from PIL import Image

    images = []
    for p in image_paths:
        img = Image.open(p)
        if img.mode != "RGB":
            img = img.convert("RGB")
        images.append(img)

    if not images:
        raise ValueError("没有可用的图片")

    first, rest = images[0], images[1:]
    first.save(str(out_path), "PDF", save_all=True, append_images=rest)


# ==================== PDF 合并 ====================
def merge_pdfs(pdf_paths: list[Path], out_path: Path):
    from pypdf import PdfWriter

    writer = PdfWriter()
    for p in pdf_paths:
        writer.append(str(p))
    with open(out_path, "wb") as f:
        writer.write(f)
    writer.close()


# ==================== PDF 拆分 ====================
def split_pdf(pdf_path: Path, out_dir: Path):
    from pypdf import PdfReader

    reader = PdfReader(str(pdf_path))
    out_dir.mkdir(parents=True, exist_ok=True)
    stem = pdf_path.stem
    outputs = []
    for i, page in enumerate(reader.pages, 1):
        w = PdfWriter() if False else None
        from pypdf import PdfWriter as W
        w = W()
        w.add_page(page)
        out = out_dir / f"{stem}_第{i}页.pdf"
        with open(out, "wb") as f:
            w.write(f)
        outputs.append(out)
    return outputs